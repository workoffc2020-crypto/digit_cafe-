"""PDF -> Image / Word / Excel conversion engine (fully offline).

Principles (same as the rest of the app):
  * Source PDFs are never modified; outputs never overwrite existing files (name (1), (2) ... is used).
  * Nothing is guessed silently: scanned pages, OCR use, low-confidence text and layout simplifications are reported.
  * Excel: IDs / phone / Aadhaar-like values stay TEXT (no leading-zero loss, no 1.23E+11), only clean numbers become numbers.
"""
from __future__ import annotations

import io
import logging
import os
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Callable, Optional

import numpy as np
from PIL import Image

try:
    import pymupdf
except ImportError:  # PyMuPDF < 1.24
    import fitz as pymupdf

from .pdf_engine import Cancelled

Image.MAX_IMAGE_PIXELS = None
log = logging.getLogger("ccs")

Progress = Optional[Callable[[int, int, str], None]]
Cancel = Optional[Callable[[], bool]]

LOW_CONF = 0.70            # OCR text below this confidence is highlighted for checking
OCR_DPI = 250
MAX_RENDER_PIXELS = 150_000_000


class PasswordRequired(Exception):
    pass


class OcrUnavailable(Exception):
    pass


# ------------------------------------------------------------------ options / report
@dataclass
class ConvertOptions:
    pages: str = ""                     # "" = all, or "1-3,5,8-"
    password: str = ""
    # images
    fmt: str = "png"                    # png / jpg / tiff
    dpi: int = 200
    jpeg_quality: int = 95
    image_mode: str = "render"          # render pages | extract embedded original images
    # OCR
    ocr: str = "auto"                   # auto (only scanned pages) / off / force
    highlight_uncertain: bool = True
    # excel
    excel_layout: str = "sheet_per_table"   # sheet_per_table / sheet_per_page / single_sheet
    text_fallback: bool = True
    convert_numbers: bool = True


@dataclass
class ConvertReport:
    kind: str
    source: str
    outputs: list = field(default_factory=list)
    pages_total: int = 0
    pages_done: int = 0
    ocr_pages: list = field(default_factory=list)
    tables: int = 0
    rows: int = 0
    text_coverage: Optional[float] = None
    warnings: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    error: str = ""

    @property
    def ok(self) -> bool:
        return not self.error and bool(self.outputs)

    def to_text(self) -> str:
        names = {"image": "PDF → Image", "word": "PDF → Word", "excel": "PDF → Excel"}
        L = [f"{names.get(self.kind, self.kind)}: {os.path.basename(self.source)}"]
        if self.error:
            L.append(f"  ✖ FAILED: {self.error}")
        else:
            L.append(f"  ✔ Pages converted: {self.pages_done}/{self.pages_total}")
            for o in self.outputs[:12]:
                L.append(f"  → {o}")
            if len(self.outputs) > 12:
                L.append(f"  … and {len(self.outputs) - 12} more files")
        if self.ocr_pages:
            L.append(f"  OCR used on page(s): {_compact_pages(self.ocr_pages)}")
        if self.kind == "excel" and not self.error:
            L.append(f"  Tables: {self.tables}   Rows written: {self.rows}")
        if self.text_coverage is not None:
            L.append(f"  Text check: {self.text_coverage * 100:.1f}% of the PDF's characters were found in the Word file")
        L += [f"  Note: {n}" for n in self.notes]
        L += [f"  ⚠ {w}" for w in self.warnings]
        return "\n".join(L)


def _compact_pages(pages) -> str:
    pages = sorted(set(pages))
    out, start, prev = [], None, None
    for p in pages:
        if start is None:
            start = prev = p
        elif p == prev + 1:
            prev = p
        else:
            out.append(f"{start}-{prev}" if start != prev else str(start))
            start = prev = p
    if start is not None:
        out.append(f"{start}-{prev}" if start != prev else str(start))
    return ", ".join(out)


# ------------------------------------------------------------------ helpers
def parse_page_range(spec: str, n: int) -> list:
    """'1-3,5,8-' -> zero-based sorted page indexes. Raises ValueError with a friendly message."""
    spec = (spec or "").strip().lower().replace(" ", "")
    if spec in ("", "all"):
        return list(range(n))
    out = []
    for part in spec.split(","):
        if not part:
            continue
        m = re.fullmatch(r"(\d*)-(\d*)", part)
        if m:
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else n
        elif part.isdigit():
            a = b = int(part)
        else:
            raise ValueError(f"Invalid page range “{part}”. Example: 1-3,5,8-")
        if a < 1 or b < a:
            raise ValueError(f"Invalid page range “{part}”.")
        if a > n:
            raise ValueError(f"Page {a} is beyond the last page ({n}).")
        out.extend(range(a - 1, min(b, n)))
    if not out:
        raise ValueError("No pages selected.")
    return sorted(set(out))


def unique_path(path: str) -> str:
    """Never overwrite: report.docx -> report (1).docx -> report (2).docx"""
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    i = 1
    while os.path.exists(f"{base} ({i}){ext}"):
        i += 1
    return f"{base} ({i}){ext}"


def _stem(path: str) -> str:
    s = os.path.splitext(os.path.basename(path))[0]
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s).strip() or "document"


def open_pdf(path: str, password: str = ""):
    doc = pymupdf.open(path)                       # raises for corrupt / unreadable files
    if not doc.is_pdf:
        doc.close()
        raise ValueError("This file is not a PDF.")
    if doc.needs_pass:
        if not password or not doc.authenticate(password):
            doc.close()
            raise PasswordRequired(os.path.basename(path))
    return doc


def page_has_text(page) -> bool:
    return len(page.get_text("text").strip()) >= 12


_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _clean(s: str) -> str:
    return _XML_BAD.sub("", s)


def _check(cancel: Cancel):
    if cancel and cancel():
        raise Cancelled()


# ------------------------------------------------------------------ OCR (local, offline)
_ocr_engine = None


def ocr_available() -> bool:
    try:
        import rapidocr_onnxruntime  # noqa: F401
        return True
    except Exception:
        return False


def _get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except Exception as e:
            raise OcrUnavailable(
                "Scanned pages need the OCR component (rapidocr-onnxruntime), which is not installed. "
                "Run:  pip install --no-deps rapidocr-onnxruntime") from e
        _ocr_engine = RapidOCR()
    return _ocr_engine


@dataclass
class Seg:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    conf: float = 1.0

    @property
    def cx(self):
        return (self.x0 + self.x1) / 2

    @property
    def cy(self):
        return (self.y0 + self.y1) / 2

    @property
    def h(self):
        return max(self.y1 - self.y0, 1e-6)


def ocr_page(page, dpi: int = OCR_DPI) -> list:
    """OCR a page; returns Seg list in PDF points (so the same grid logic works for text and scanned pages)."""
    pix = page.get_pixmap(dpi=dpi, alpha=False, colorspace=pymupdf.csRGB)
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    img = np.ascontiguousarray(img[..., ::-1])      # RGB -> BGR (what the OCR engine expects)
    result, _ = _get_ocr()(img)
    k = 72.0 / dpi
    segs = []
    for box, txt, conf in (result or []):
        txt = _clean(str(txt)).strip()
        if not txt:
            continue
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        segs.append(Seg(txt, min(xs) * k, min(ys) * k, max(xs) * k, max(ys) * k, float(conf)))
    return segs


# ------------------------------------------------------------------ layout helpers (rows / columns)
def rows_from_segments(segs: list) -> list:
    """Group segments into visual rows (top to bottom), each row sorted left to right."""
    if not segs:
        return []
    med_h = float(np.median([s.h for s in segs]))
    segs = sorted(segs, key=lambda s: s.cy)
    rows, cur, cur_cy = [], [], None
    for s in segs:
        if cur and abs(s.cy - cur_cy) <= 0.6 * med_h:
            cur.append(s)
            cur_cy = float(np.mean([t.cy for t in cur]))
        else:
            if cur:
                rows.append(sorted(cur, key=lambda t: t.x0))
            cur, cur_cy = [s], s.cy
    if cur:
        rows.append(sorted(cur, key=lambda t: t.x0))
    return rows


def words_to_segments(words: list) -> list:
    """PyMuPDF words -> phrase segments (neighbouring words joined unless a column-sized gap separates them)."""
    ws = [Seg(_clean(w[4]), w[0], w[1], w[2], w[3]) for w in words if _clean(w[4]).strip()]
    if not ws:
        return []
    med_h = float(np.median([w.h for w in ws]))
    gap = max(6.0, 0.9 * med_h)
    out = []
    for row in rows_from_segments(ws):
        cur = row[0]
        for w in row[1:]:
            if w.x0 - cur.x1 <= gap:
                cur = Seg(cur.text + " " + w.text, cur.x0, min(cur.y0, w.y0), w.x1, max(cur.y1, w.y1))
            else:
                out.append(cur)
                cur = w
        out.append(cur)
    return out


@dataclass
class Cell:
    text: str
    low: bool = False


def grid_from_segments(segs: list) -> list:
    """Segments -> rows x columns grid. Column bands come from rows that have 2+ segments; other rows go to column A."""
    rows = rows_from_segments(segs)
    if not rows:
        return []
    med_h = float(np.median([s.h for s in segs]))
    tol = 0.5 * med_h
    intervals = sorted((s.x0, s.x1) for r in rows if len(r) >= 2 for s in r)
    bands = []
    for a, b in intervals:
        if bands and a <= bands[-1][1] + tol:
            bands[-1][1] = max(bands[-1][1], b)
        else:
            bands.append([a, b])
    if not bands:
        bands = [[min(s.x0 for s in segs), max(s.x1 for s in segs)]]
    grid = []
    for r in rows:
        line = [Cell("") for _ in bands]
        if len(r) == 1:
            line[0] = Cell(r[0].text, r[0].conf < LOW_CONF)
        else:
            for s in r:
                j = min(range(len(bands)), key=lambda k: 0 if bands[k][0] <= s.cx <= bands[k][1]
                        else min(abs(s.cx - bands[k][0]), abs(s.cx - bands[k][1])))
                c = line[j]
                line[j] = Cell((c.text + " " + s.text).strip(), c.low or s.conf < LOW_CONF)
        grid.append(line)
    return _trim_grid(grid)


def _trim_grid(grid: list) -> list:
    grid = [r for r in grid if any(c.text.strip() for c in r)]
    if not grid:
        return []
    keep = [j for j in range(len(grid[0])) if any(r[j].text.strip() for r in grid)]
    return [[r[j] for j in keep] for r in grid]


# ------------------------------------------------------------------ Excel value rules
_NUM_BODY = re.compile(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d{1,2}(?:,\d{2})+,\d{3}(?:\.\d+)?|\d+(?:\.\d+)?")
_CUR = {"₹": "₹", "$": "$", "€": "€", "£": "£"}


def smart_value(text: str):
    """Return (value, number_format). Only clean numbers are converted; identifiers stay text.

    Kept as text: leading zeros ("0123"), 10+ digit plain numbers (phone / Aadhaar / account), >15 digits,
    dates, anything with letters. Handles 1,234.50 / 12,34,567 / (1,234) / -5 / 12% / ₹1,250.
    """
    t = (text or "").strip()
    if not t:
        return "", None
    u, neg, pct, cur = t, False, False, ""
    if u.startswith("(") and u.endswith(")"):
        neg, u = True, u[1:-1].strip()
    if u.startswith("-") or u.startswith("−"):
        neg, u = True, u[1:].strip()
    m = re.match(r"^(Rs\.?|INR)\s?", u, re.I)
    if m:
        cur, u = "₹", u[m.end():]
    elif u[:1] in _CUR:
        cur, u = _CUR[u[0]], u[1:].strip()
    if u.endswith("%"):
        pct, u = True, u[:-1].strip()
    if not _NUM_BODY.fullmatch(u):
        return t, None
    plain = u.replace(",", "")
    ip, _, fp = plain.partition(".")
    if len(ip) > 1 and ip.startswith("0"):
        return t, None                                  # 0123 -> keep as text
    if "," not in u and not fp and len(ip) >= 10:
        return t, None                                  # phone / Aadhaar / account number
    if len(ip) + len(fp) > 15:
        return t, None
    val = float(plain) if fp else int(plain)
    if neg:
        val = -val
    d = len(fp)
    if pct:
        val = val / 100
        fmt = "0" + ("." + "0" * d if d else "") + "%"
    else:
        fmt = ("#,##0" if "," in u else "0") + ("." + "0" * d if d else "")
        if cur:
            fmt = f'"{cur}"{fmt}'
    return val, fmt


# ------------------------------------------------------------------ PDF -> Image
def pdf_to_images(pdf: str, out_dir: str, opts: ConvertOptions, progress: Progress = None, cancel: Cancel = None) -> ConvertReport:
    rep = ConvertReport("image", pdf)
    doc = open_pdf(pdf, opts.password)
    try:
        idxs = parse_page_range(opts.pages, doc.page_count)
        rep.pages_total = len(idxs)
        folder = unique_path(os.path.join(out_dir, _stem(pdf) + ("_images" if opts.image_mode == "extract" else "_pages")))
        os.makedirs(folder, exist_ok=True)
        if opts.image_mode == "extract":
            _extract_images(doc, idxs, folder, rep, progress, cancel)
        else:
            _render_pages(doc, idxs, folder, opts, rep, progress, cancel)
        if not rep.outputs:
            try:
                os.rmdir(folder)
            except OSError:
                pass
    finally:
        doc.close()
    return rep


def _render_pages(doc, idxs, folder, opts, rep, progress, cancel):
    fmt = {"jpg": "jpg", "jpeg": "jpg", "tif": "tiff", "tiff": "tiff"}.get(opts.fmt.lower(), "png")
    dpi = int(min(max(opts.dpi, 50), 600))
    width = len(str(len(idxs)))
    for n, i in enumerate(idxs):
        _check(cancel)
        if progress:
            progress(n, len(idxs), f"Rendering page {i + 1}")
        page = doc[i]
        d = dpi
        px = (page.rect.width / 72 * d) * (page.rect.height / 72 * d)
        if px > MAX_RENDER_PIXELS:
            d = int(d * (MAX_RENDER_PIXELS / px) ** 0.5)
            rep.warnings.append(f"Page {i + 1} is very large; rendered at {d} DPI instead of {dpi} to save memory.")
        pix = page.get_pixmap(dpi=d, alpha=False, colorspace=pymupdf.csRGB)
        im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        path = os.path.join(folder, f"{_stem(rep.source)}_page_{str(i + 1).zfill(max(width, 3))}.{'jpg' if fmt == 'jpg' else fmt}")
        if fmt == "jpg":
            im.save(path, "JPEG", quality=int(opts.jpeg_quality), subsampling=0 if opts.jpeg_quality >= 90 else 2,
                    optimize=True, dpi=(d, d))
        elif fmt == "tiff":
            im.save(path, "TIFF", compression="tiff_lzw", dpi=(d, d))
        else:
            im.save(path, "PNG", compress_level=6, dpi=(d, d))
        rep.outputs.append(path)
        rep.pages_done += 1
    if fmt == "jpg":
        rep.notes.append(f"JPG is lossy (quality {opts.jpeg_quality}). Choose PNG or TIFF for lossless pages.")
    else:
        rep.notes.append(f"{fmt.upper()} output is lossless at {dpi} DPI.")
    if progress:
        progress(len(idxs), len(idxs), "Done")


def _extract_images(doc, idxs, folder, rep, progress, cancel):
    seen, count, smask_warn = set(), 0, False
    for n, i in enumerate(idxs):
        _check(cancel)
        if progress:
            progress(n, len(idxs), f"Extracting images from page {i + 1}")
        k = 0
        for img in doc[i].get_images(full=True):
            xref, smask = img[0], img[1]
            if xref in seen:
                continue
            seen.add(xref)
            try:
                info = doc.extract_image(xref)
            except Exception as e:
                rep.warnings.append(f"Page {i + 1}: image could not be extracted ({e}).")
                continue
            if not info or not info.get("image"):
                continue
            k += 1
            path = os.path.join(folder, f"{_stem(rep.source)}_p{i + 1:03d}_img{k}.{info.get('ext', 'png')}")
            with open(path, "wb") as f:
                f.write(info["image"])
            rep.outputs.append(path)
            count += 1
            if smask:
                smask_warn = True
        rep.pages_done += 1
    if count == 0:
        rep.warnings.append("No embedded images found. Use “Render pages” to turn pages into pictures.")
    else:
        rep.notes.append("Original embedded image data was saved byte-for-byte (JPEG stays the exact original JPEG; no re-encoding).")
    if smask_warn:
        rep.warnings.append("Some images have transparency masks that are not included in the extracted files.")
    if progress:
        progress(len(idxs), len(idxs), "Done")


# ------------------------------------------------------------------ PDF -> Word
def pdf_to_word(pdf: str, out_docx: str, opts: ConvertOptions, progress: Progress = None, cancel: Cancel = None) -> ConvertReport:
    rep = ConvertReport("word", pdf)
    doc = open_pdf(pdf, opts.password)
    try:
        idxs = parse_page_range(opts.pages, doc.page_count)
        rep.pages_total = len(idxs)
        has_text = [page_has_text(doc[i]) for i in idxs]
        if opts.ocr == "force":
            need_ocr = [True] * len(idxs)
        elif opts.ocr == "off":
            need_ocr = [False] * len(idxs)
        else:
            need_ocr = [not t for t in has_text]
        out = unique_path(out_docx)
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)

        if not any(need_ocr):
            blank = [i + 1 for i, t in zip(idxs, has_text) if not t]
            used_own = False
            _check(cancel)
            if progress:
                progress(0, 1, "Converting layout (this can take a while for big files)…")
            try:
                _pdf2docx(pdf, out, idxs, opts.password, cancel)
            except Cancelled:
                raise
            except Exception as e:
                log.exception("pdf2docx failed; using simple text builder")
                rep.warnings.append(f"Layout converter failed ({e}); a simplified text version was created instead.")
                _build_docx_own(doc, idxs, [False] * len(idxs), out, opts, rep, progress, cancel)
                used_own = True
            if blank:
                rep.warnings.append(f"Page(s) {_compact_pages(blank)} contain no text (scanned images). They were inserted as pictures; "
                                    "turn OCR on to get editable text.")
            if not used_own:
                rep.notes.append("Layout, fonts, tables and pictures were converted with the layout engine. "
                                 "Complex designs can still shift slightly in Word; please review the result.")
        else:
            if not all(need_ocr):
                rep.warnings.append("Mixed text and scanned pages: a simplified layout was used (text formatting kept; "
                                    "tables and pictures on text pages are not reproduced).")
            _build_docx_own(doc, idxs, need_ocr, out, opts, rep, progress, cancel)

        rep.outputs.append(out)
        rep.pages_done = len(idxs)
        rep.text_coverage = _word_coverage(doc, idxs, [not n for n in need_ocr], out)
        if rep.text_coverage is not None and rep.text_coverage < 0.97:
            rep.warnings.append("Less than 97% of the PDF text was found in the Word file. Please compare both before use.")
        if rep.ocr_pages:
            rep.warnings.append("OCR text can contain mistakes. Uncertain words are highlighted in yellow. "
                                "Always verify names, dates and numbers.")
        if progress:
            progress(1, 1, "Done")
    finally:
        doc.close()
    return rep


def _pdf2docx(pdf, out, idxs, password, cancel):
    from pdf2docx import Converter
    logging.disable(logging.INFO)
    try:
        cv = Converter(pdf, password=password or None)
        try:
            cv.convert(out, pages=list(idxs))
        finally:
            cv.close()
    finally:
        logging.disable(logging.NOTSET)
    _check(cancel)


def _ocr_body_size(segs: list) -> float:
    """Typical font size of a scanned page (pt). Estimated from text WIDTH per character: box heights are
    unreliable (short words get padded boxes), long lines are not."""
    est = [(s.x1 - s.x0) / len(s.text) / 0.5 for s in segs if len(s.text) >= 8]
    if est:
        return float(np.median(est))
    return max(float(np.median([s.h for s in segs])) * 0.8, 8.0) if segs else 11.0


def _ocr_seg_size(s: Seg, base: float) -> float:
    if len(s.text) >= 8:
        est = (s.x1 - s.x0) / len(s.text) / 0.5
        if est > 1.3 * base:                         # clearly larger -> heading
            return est
    return base


def _build_docx_own(doc, idxs, need_ocr, out, opts, rep, progress, cancel):
    import docx
    from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
    from docx.shared import Inches, Pt, RGBColor

    d = docx.Document()
    first = doc[idxs[0]].rect
    sec = d.sections[0]
    sec.page_width, sec.page_height = Pt(first.width), Pt(first.height)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.7))
    omitted_images = False
    for n, i in enumerate(idxs):
        _check(cancel)
        page = doc[i]
        if progress:
            progress(n, len(idxs), f"{'OCR' if need_ocr[n] else 'Reading'} page {i + 1}")
        if need_ocr[n]:
            rep.ocr_pages.append(i + 1)
            segs = ocr_page(page)
            if not segs:
                rep.warnings.append(f"Page {i + 1}: OCR found no text.")
            base_size = _ocr_body_size(segs)
            for row in rows_from_segments(segs):
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                med_h = max(float(np.median([s.h for s in row])), 1.0)
                for k, s in enumerate(row):
                    if k:
                        p.add_run("\t" if s.x0 - row[k - 1].x1 > 0.8 * med_h else " ")   # separate OCR boxes + a real gap = column
                    r = p.add_run(s.text)
                    r.font.size = Pt(round(min(max(_ocr_seg_size(s, base_size), 8), 26) * 2) / 2)
                    if opts.highlight_uncertain and s.conf < LOW_CONF:
                        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif page_has_text(page):
            for b in page.get_text("dict")["blocks"]:
                if b.get("type") == 1:
                    omitted_images = True
                    continue
                if b.get("type") != 0:
                    continue
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                lines = b.get("lines", [])
                for li, line in enumerate(lines):
                    for sp in line.get("spans", []):
                        t = _clean(sp.get("text", ""))
                        if not t:
                            continue
                        r = p.add_run(t)
                        r.font.size = Pt(min(max(sp.get("size", 11), 6), 48))
                        r.bold = bool(sp.get("flags", 0) & 16)
                        r.italic = bool(sp.get("flags", 0) & 2)
                        col = int(sp.get("color", 0))
                        r.font.color.rgb = RGBColor((col >> 16) & 255, (col >> 8) & 255, col & 255)
                    if li < len(lines) - 1:
                        p.add_run(" ")
        else:   # no text and OCR off: keep the page as a picture so no content is lost
            pix = page.get_pixmap(dpi=150, alpha=False, colorspace=pymupdf.csRGB)
            d.add_picture(io.BytesIO(pix.tobytes("png")), width=Pt(first.width - 100))
        if n < len(idxs) - 1:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    if omitted_images:
        rep.warnings.append("Pictures on text pages were not copied in simplified mode.")
    d.save(out)


def _norm(s: str) -> str:
    return re.sub(r"\W+", "", s, flags=re.UNICODE).lower()


def _word_coverage(doc, idxs, text_pages, out) -> Optional[float]:
    """Fraction of the PDF's text characters that appear in the .docx (only for pages that had real text)."""
    try:
        import docx
        src = "".join(_norm(doc[i].get_text("text")) for i, t in zip(idxs, text_pages) if t)
        if len(src) < 20:
            return None
        d = docx.Document(out)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.append(c.text)
        dst = _norm(" ".join(parts))
        a, b = Counter(src), Counter(dst)
        return sum(min(v, b.get(k, 0)) for k, v in a.items()) / max(sum(a.values()), 1)
    except Exception:
        log.exception("coverage check failed")
        return None


# ------------------------------------------------------------------ PDF -> Excel
def _table_grid(page, table, words) -> list:
    """Cell text taken from the words that sit inside each cell (robust: avoids garbled table.extract() text)."""
    grid = []
    for r in table.rows:
        line = []
        for c in r.cells:
            if c is None:
                line.append(Cell(""))
                continue
            x0, y0, x1, y1 = c
            ws = [w for w in words if x0 <= (w[0] + w[2]) / 2 <= x1 and y0 <= (w[1] + w[3]) / 2 <= y1]
            ws.sort(key=lambda w: (round(w[1] / 3), w[0]))
            line.append(Cell(_clean(" ".join(w[4] for w in ws)).strip()))
        grid.append(line)
    return _trim_grid(grid)


def pdf_to_excel(pdf: str, out_xlsx: str, opts: ConvertOptions, progress: Progress = None, cancel: Cancel = None) -> ConvertReport:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    rep = ConvertReport("excel", pdf)
    doc = open_pdf(pdf, opts.password)
    try:
        idxs = parse_page_range(opts.pages, doc.page_count)
        rep.pages_total = len(idxs)
        blocks = []                                     # (page_no, label, grid, is_table)
        for n, i in enumerate(idxs):
            _check(cancel)
            page = doc[i]
            scanned = (opts.ocr == "force") or (opts.ocr == "auto" and not page_has_text(page))
            if progress:
                progress(n, len(idxs), f"{'OCR' if scanned else 'Reading tables on'} page {i + 1}")
            if scanned:
                if opts.ocr == "off":
                    continue
                rep.ocr_pages.append(i + 1)
                grid = grid_from_segments(ocr_page(page))
                if grid:
                    blocks.append((i + 1, f"P{i + 1}", grid, False))
                else:
                    rep.warnings.append(f"Page {i + 1}: OCR found no text.")
                rep.pages_done += 1
                continue
            if opts.ocr == "off" and not page_has_text(page):
                rep.warnings.append(f"Page {i + 1} has no text (scanned image). Turn OCR on to read it.")
                rep.pages_done += 1
                continue
            words = page.get_text("words")
            found = 0
            try:
                tables = page.find_tables().tables
            except Exception as e:
                tables = []
                rep.warnings.append(f"Page {i + 1}: table detection failed ({e}).")
            for t in tables:
                g = _table_grid(page, t, words)
                if g:
                    found += 1
                    blocks.append((i + 1, f"P{i + 1}_T{found}", g, True))
                    rep.tables += 1
            if not found:
                if opts.text_fallback:
                    g = grid_from_segments(words_to_segments(words))
                    if g:
                        blocks.append((i + 1, f"P{i + 1}", g, False))
                        rep.warnings.append(f"Page {i + 1}: no table lines detected; text was arranged into rows and columns by position. Please check it.")
                else:
                    rep.warnings.append(f"Page {i + 1}: no tables found.")
            rep.pages_done += 1
        if not blocks:
            rep.error = "No tables or text could be extracted from the selected pages."
            return rep

        wb = Workbook()
        wb.remove(wb.active)
        used_names = set()

        def sheet(name):
            name = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
            base, k = name, 1
            while name.lower() in used_names:
                k += 1
                name = f"{base[:28]}_{k}"
            used_names.add(name.lower())
            return wb.create_sheet(name)

        yellow = PatternFill("solid", start_color="FFF59D")
        bold = Font(bold=True)

        def write(ws, grid, row0, header):
            for r, line in enumerate(grid):
                for c, cell in enumerate(line):
                    if not cell.text:
                        continue
                    val, fmt = smart_value(cell.text) if opts.convert_numbers else (cell.text, None)
                    x = ws.cell(row=row0 + r, column=c + 1, value=val)
                    if fmt:
                        x.number_format = fmt
                    elif isinstance(val, str):
                        x.number_format = "@"             # keep as text, even if Excel would reinterpret it
                    if cell.low:
                        x.fill = yellow
                    if header and r == 0:
                        x.font = bold
            return row0 + len(grid)

        layout = opts.excel_layout
        if layout == "sheet_per_table":
            for pno, label, grid, is_table in blocks:
                ws = sheet(label)
                write(ws, grid, 1, is_table)
                if is_table:
                    ws.freeze_panes = "A2"
                rep.rows += len(grid)
        else:
            sheets = {}
            for pno, label, grid, is_table in blocks:
                key = "Data" if layout == "single_sheet" else f"Page {pno}"
                ws = sheets.get(key) or sheets.setdefault(key, sheet(key))
                row0 = ws.max_row + 2 if ws.max_row > 1 or ws.cell(1, 1).value is not None else 1
                if layout == "single_sheet" and len(idxs) > 1:
                    ws.cell(row=row0, column=1, value=f"— Page {pno} —").font = bold
                    row0 += 1
                write(ws, grid, row0, is_table)
                rep.rows += len(grid)
        for ws in wb.worksheets:
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 60)

        out = unique_path(out_xlsx)
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        wb.save(out)
        rep.outputs.append(out)
        rep.notes.append("Numbers like 1,234.50 became real numbers. IDs, phone/Aadhaar numbers and values with leading zeros stay text. "
                         "Dates are kept as text on purpose.")
        if rep.ocr_pages:
            rep.warnings.append("OCR text can contain mistakes. Uncertain cells are highlighted in yellow. Always verify numbers.")
        if progress:
            progress(len(idxs), len(idxs), "Done")
    finally:
        doc.close()
    return rep


# ------------------------------------------------------------------ batch runner (used by the GUI)
def run_batch(kind: str, files: list, out_dir: str, opts: ConvertOptions, progress: Progress = None, cancel: Cancel = None) -> list:
    """Convert many PDFs. One bad file never stops the others; every file gets its own report."""
    os.makedirs(out_dir, exist_ok=True)
    reports = []
    for fi, f in enumerate(files):
        _check(cancel)
        name = os.path.basename(f)

        def sub(a, b, m, _fi=fi, _name=name):
            if progress:
                progress(_fi * 1000 + int(1000 * a / max(b, 1)), len(files) * 1000, f"{_name}: {m}")

        try:
            if kind == "image":
                rep = pdf_to_images(f, out_dir, opts, sub, cancel)
            elif kind == "word":
                rep = pdf_to_word(f, os.path.join(out_dir, _stem(f) + ".docx"), opts, sub, cancel)
            elif kind == "excel":
                rep = pdf_to_excel(f, os.path.join(out_dir, _stem(f) + ".xlsx"), opts, sub, cancel)
            else:
                raise ValueError(f"Unknown conversion: {kind}")
        except Cancelled:
            raise
        except PasswordRequired:
            rep = ConvertReport(kind, f, error="This PDF is password protected. Enter the password and try again.")
        except OcrUnavailable as e:
            rep = ConvertReport(kind, f, error=str(e))
        except Exception as e:
            log.exception("PDF conversion failed: %s", f)
            rep = ConvertReport(kind, f, error=f"{type(e).__name__}: {e}")
        reports.append(rep)
    if progress:
        progress(len(files) * 1000, len(files) * 1000, "Done")
    return reports
