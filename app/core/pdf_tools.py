"""PDF -> Images / Word / Excel. 100% offline.

Libraries (permissive licences): pypdfium2 (page rendering), pdfplumber/pdfminer.six (text, layout, tables),
python-docx, openpyxl, pikepdf (original embedded images).

Honest limits (shown to the user in the app too):
  * Word/Excel need a *text layer*. Scanned PDFs (pictures of pages) have none; OCR is NOT included.
    For Word, scanned pages are inserted as pictures so nothing is lost; for Excel they are reported and skipped.
  * Word output re-creates text, tables and pictures; it is not a pixel-perfect copy of complex layouts.
  * Excel tables are detected from ruling lines first, then from text alignment (flagged "verify").
Source PDFs are never modified.
"""
from __future__ import annotations

import io
import logging
import os
import re
import warnings
from collections import Counter
from dataclasses import dataclass, field
from statistics import median
from typing import Callable, Optional

from PIL import Image

from .pdf_engine import Cancelled

Image.MAX_IMAGE_PIXELS = None
log = logging.getLogger("ccs")

MAX_RENDER_SIDE = 16000          # px; protects RAM on huge pages / very high DPI


class PdfToolError(Exception):
    """A problem with a PDF that can be explained to the user."""


class PdfPasswordError(PdfToolError):
    """The PDF is password protected and no / a wrong password was given."""


@dataclass
class ToolResult:
    source: str
    outputs: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    stats: dict = field(default_factory=dict)
    error: str = ""

    def summary(self) -> str:
        name = os.path.basename(self.source)
        if self.error:
            return f"✖ {name}: {self.error}"
        lines = [f"✔ {name}"]
        lines += [f"    → {o}" for o in self.outputs]
        if self.stats:
            lines.append("    " + ", ".join(f"{k}: {v}" for k, v in self.stats.items()))
        lines += [f"    note: {n}" for n in self.notes]
        lines += [f"    ⚠ {w}" for w in self.warnings]
        return "\n".join(lines)


# ------------------------------------------------------------------ helpers
def parse_pages(spec: str, n: int) -> list:
    """'1-3,5,8-' -> zero-based page indexes. Empty / 'all' = every page."""
    spec = (spec or "").strip().lower()
    if spec in ("", "all", "*"):
        return list(range(n))
    out = []
    for tok in re.split(r"[,\s;]+", spec):
        if not tok:
            continue
        m = re.fullmatch(r"(\d*)\s*-\s*(\d*)", tok)
        if m and (m.group(1) or m.group(2)):
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else n
        elif tok.isdigit():
            a = b = int(tok)
        else:
            raise PdfToolError(f"Cannot understand page range “{tok}”. Use for example 1-3,5,8-")
        if a < 1 or b < a or b > n:
            raise PdfToolError(f"Page range {tok} is outside this PDF (it has {n} pages).")
        out.extend(range(a - 1, b))
    seen, res = set(), []
    for i in out:
        if i not in seen:
            seen.add(i)
            res.append(i)
    return res


def unique_path(path: str) -> str:
    """Never overwrite an existing file or folder: name.ext -> name (1).ext"""
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    i = 1
    while os.path.exists(f"{base} ({i}){ext}"):
        i += 1
    return f"{base} ({i}){ext}"


def _open_pikepdf(path: str, password: Optional[str]):
    import pikepdf
    try:
        return pikepdf.open(path, password=password or "")
    except pikepdf.PasswordError:
        raise PdfPasswordError("This PDF is password protected" + (" (wrong password)." if password else "."))
    except Exception as e:
        raise PdfToolError(f"Cannot open this PDF (damaged or not a PDF): {e}")


def needs_password(path: str) -> bool:
    import pikepdf
    try:
        pikepdf.open(path).close()
        return False
    except pikepdf.PasswordError:
        return True
    except Exception:
        return False


def page_count(path: str, password: Optional[str] = None) -> int:
    with _open_pikepdf(path, password) as pdf:
        return len(pdf.pages)


def _chk(cancel):
    if cancel and cancel():
        raise Cancelled()


# ------------------------------------------------------------------ PDF -> images
def pdf_to_images(pdf: str, out_dir: str, fmt: str = "png", dpi: int = 200, pages: str = "",
                  password: Optional[str] = None, mode: str = "render", jpeg_quality: int = 95,
                  progress=None, cancel=None) -> ToolResult:
    """mode='render': every page as a picture. mode='extract': the original embedded pictures, byte-for-byte."""
    res = ToolResult(source=pdf)
    n = page_count(pdf, password)
    idxs = parse_pages(pages, n)
    stem = os.path.splitext(os.path.basename(pdf))[0]
    folder = unique_path(os.path.join(out_dir, f"{stem}_images"))
    os.makedirs(folder, exist_ok=True)
    fmt = fmt.lower().replace("jpeg", "jpg")

    if mode == "extract":
        _extract_embedded(pdf, folder, idxs, password, res, progress, cancel)
    else:
        import pypdfium2 as pdfium
        try:
            doc = pdfium.PdfDocument(pdf, password=password or None)
        except Exception as e:
            raise PdfToolError(f"Cannot render this PDF: {e}")
        clamped = False
        try:
            for k, i in enumerate(idxs):
                _chk(cancel)
                page = doc[i]
                w, h = page.get_size()
                scale = dpi / 72.0
                if max(w, h) * scale > MAX_RENDER_SIDE:
                    scale = MAX_RENDER_SIDE / max(w, h)
                    clamped = True
                im = page.render(scale=scale).to_pil()
                page.close()
                real_dpi = round(scale * 72)
                path = os.path.join(folder, f"{stem}_page_{i + 1:03d}.{fmt}")
                if fmt == "jpg":
                    im.convert("RGB").save(path, "JPEG", quality=jpeg_quality, subsampling=0, dpi=(real_dpi, real_dpi))
                elif fmt == "tif":
                    im.convert("RGB").save(path, "TIFF", compression="tiff_lzw", dpi=(real_dpi, real_dpi))
                else:
                    im.convert("RGB").save(path, "PNG", compress_level=6, dpi=(real_dpi, real_dpi))
                res.outputs.append(path)
                if progress:
                    progress(k + 1, len(idxs), f"Page {i + 1} → image")
        finally:
            doc.close()
        if clamped:
            res.warnings.append(f"Some pages were rendered below {dpi} DPI to stay under {MAX_RENDER_SIDE}px per side.")
        res.notes.append(f"Rendered at {dpi} DPI; {fmt.upper()} output.")
    res.stats = {"pages": len(idxs), "files": len(res.outputs)}
    res.notes.insert(0, f"Folder: {folder}")
    if not res.outputs:
        res.warnings.append("No images were produced.")
    return res


def _extract_embedded(pdf, folder, idxs, password, res, progress, cancel):
    import pikepdf
    stem = os.path.splitext(os.path.basename(pdf))[0]
    seen = set()
    skipped = 0
    with _open_pikepdf(pdf, password) as doc:
        for k, i in enumerate(idxs):
            _chk(cancel)
            page = doc.pages[i]
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                items = list(page.images.items())
            for j, (_, obj) in enumerate(items, 1):
                key = obj.objgen
                if key in seen:                      # same picture (e.g. a logo) repeated on many pages
                    continue
                seen.add(key)
                try:
                    if int(obj.get("/Width", 0)) < 16 or int(obj.get("/Height", 0)) < 16:
                        continue
                    ext = pikepdf.PdfImage(obj).extract_to(fileprefix=os.path.join(folder, f"{stem}_p{i + 1:03d}_img{j:02d}"))
                    res.outputs.append(ext)
                except Exception as e:
                    skipped += 1
                    log.warning("embedded image skipped: %s", e)
            if progress:
                progress(k + 1, len(idxs), f"Page {i + 1} → embedded pictures")
    res.notes.append("Original embedded pictures extracted without re-encoding (JPEG stays the exact same file).")
    if skipped:
        res.warnings.append(f"{skipped} embedded picture(s) used an unsupported format and were skipped.")
    if not res.outputs:
        res.warnings.append("No embedded pictures found. Use “Render pages” to get each page as an image.")


# ------------------------------------------------------------------ pdfplumber helpers
def _plumber(pdf: str, password: Optional[str]):
    import pdfplumber
    _open_pikepdf(pdf, password).close()          # clear, friendly errors for bad / protected files first
    try:
        return pdfplumber.open(pdf, password=password or None)
    except Exception as e:
        raise PdfToolError(f"Cannot read this PDF: {e}")


def _free(page):
    try:
        page.flush_cache()
    except Exception:
        pass
    try:
        page.close()
    except Exception:
        pass


def _in_boxes(obj, boxes) -> bool:
    cx = (obj["x0"] + obj["x1"]) / 2
    cy = (obj["top"] + obj["bottom"]) / 2
    return any(b[0] - 1 <= cx <= b[2] + 1 and b[1] - 1 <= cy <= b[3] + 1 for b in boxes)


def _clean_rows(rows) -> list:
    out = [[(c or "").replace("\r", "").strip() for c in r] for r in rows or []]
    out = [r for r in out if any(r)]
    if not out:
        return []
    ncol = max(len(r) for r in out)
    out = [r + [""] * (ncol - len(r)) for r in out]
    keep = [j for j in range(ncol) if any(r[j] for r in out)]
    return [[r[j] for j in keep] for r in out]


def _ruled_tables(page):
    """Tables drawn with ruling lines. Returns [(bbox, rows)], ignoring decorative single boxes."""
    found = []
    try:
        tables = page.find_tables()
    except Exception:
        return found
    for t in tables:
        rows = _clean_rows(t.extract())
        if len(rows) >= 2 and len(rows[0]) >= 2 and sum(1 for r in rows for c in r if c) >= 3:
            found.append((t.bbox, rows))
    return found


def _is_scan(page, text_chars: int) -> bool:
    return text_chars < 10 and len(page.images) > 0


# ------------------------------------------------------------------ PDF -> Word
def _is_bold(fontname: str) -> bool:
    f = (fontname or "").lower()
    return any(k in f for k in ("bold", "black", "heavy", "demi"))


def _is_italic(fontname: str) -> bool:
    f = (fontname or "").lower()
    return "italic" in f or "oblique" in f


def _body_font_size(plumber_pdf, idxs) -> float:
    sizes = Counter()
    for i in idxs[:6]:
        pg = plumber_pdf.pages[i]
        for c in pg.chars:
            if c["text"].strip():
                sizes[round(c["size"], 1)] += 1
        _free(pg)
    return sizes.most_common(1)[0][0] if sizes else 11.0


def _paragraphs(page, boxes, page_w):
    pg = page.filter(lambda o: not (o["object_type"] == "char" and _in_boxes(o, boxes))) if boxes else page
    lines = pg.extract_text_lines(return_chars=True, strip=True) or []
    lines = [l for l in lines if l["text"].strip()]
    if not lines:
        return []
    right_edge = max(l["x1"] for l in lines)
    left_edge = min(l["x0"] for l in lines)
    paras = []
    for L in lines:
        chars = [c for c in L["chars"] if c["text"].strip()]
        if not chars:
            continue
        size = float(median(c["size"] for c in chars))
        bold = sum(_is_bold(c["fontname"]) for c in chars) / len(chars) > 0.6
        italic = sum(_is_italic(c["fontname"]) for c in chars) / len(chars) > 0.6
        cur = paras[-1] if paras else None
        lh = L["bottom"] - L["top"]
        cont = False
        if cur is not None:
            gap = L["top"] - cur["bottom"]
            full_prev = cur["last_x1"] >= right_edge - 0.06 * page_w     # previous line reached the right margin
            cont = (full_prev and gap <= 0.5 * cur["line_h"] and abs(size - cur["size"]) <= 0.6
                    and bold == cur["bold"] and abs(L["x0"] - cur["x0"]) <= 3 * size)
        if cont:
            cur["text"] += " " + L["text"]
            cur["bottom"] = L["bottom"]
            cur["x1"] = max(cur["x1"], L["x1"])
            cur["last_x1"] = L["x1"]
        else:
            paras.append({"text": L["text"], "top": L["top"], "bottom": L["bottom"], "x0": L["x0"], "x1": L["x1"],
                          "last_x1": L["x1"], "size": size, "bold": bold, "italic": italic, "line_h": lh,
                          "gap": (L["top"] - cur["bottom"]) if cur is not None else None, "left_edge": left_edge})
    return paras


def pdf_to_word(pdf: str, out_path: str, pages: str = "", password: Optional[str] = None,
                keep_images: bool = True, progress=None, cancel=None) -> ToolResult:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    res = ToolResult(source=pdf)
    n = page_count(pdf, password)
    idxs = parse_pages(pages, n)
    stats = Counter()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    margin = 0.6

    with _plumber(pdf, password) as plumber:
        body = _body_font_size(plumber, idxs)
        first_size = None
        for k, i in enumerate(idxs):
            _chk(cancel)
            page = plumber.pages[i]
            pw, ph = float(page.width), float(page.height)
            if first_size is None:
                sec = doc.sections[0]
                first_size = (pw, ph)
            elif (round(pw), round(ph)) != (round(first_size[0]), round(first_size[1])):
                sec = doc.add_section(WD_SECTION.NEW_PAGE)
                first_size = (pw, ph)
            else:
                doc.add_page_break()
                sec = None
            if sec is not None:
                sec.page_width, sec.page_height = Pt(pw), Pt(ph)
                sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(margin)
            avail_w = pw / 72.0 - 2 * margin

            text = page.extract_text() or ""
            nchar = len(text.strip())
            garbled = nchar > 20 and text.count("(cid:") * 6 > nchar * 0.10
            if nchar < 10 or garbled:
                if len(page.images) or garbled:
                    im = page.to_image(resolution=170).original.convert("RGB")
                    buf = io.BytesIO()
                    im.save(buf, "JPEG", quality=88)
                    doc.add_picture(io.BytesIO(buf.getvalue()), width=Inches(min(avail_w, pw / 72.0)))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    stats["pictures"] += 1
                    res.warnings.append(
                        f"Page {i + 1}: " + ("text could not be decoded" if garbled else "scanned page, no text layer")
                        + " — inserted as a picture (OCR is not included, so it is not editable).")
                _free(page)
                if progress:
                    progress(k + 1, len(idxs), f"Page {i + 1}")
                continue

            blocks = []
            tables = _ruled_tables(page)
            boxes = [t[0] for t in tables]
            for bbox, rows in tables:
                blocks.append((bbox[1], "t", rows))
            for p in _paragraphs(page, boxes, pw):
                blocks.append((p["top"], "p", p))
            if keep_images:
                for img in page.images:
                    x0, x1 = max(img["x0"], 0), min(img["x1"], pw)
                    t0, b1 = max(img["top"], 0), min(img["bottom"], ph)
                    w, h = x1 - x0, b1 - t0
                    if w < 24 or h < 24:
                        continue
                    if w * h > 0.85 * pw * ph:
                        res.notes.append(f"Page {i + 1}: full-page background picture skipped.")
                        continue
                    try:
                        im = page.crop((x0, t0, x1, b1)).to_image(resolution=150).original.convert("RGB")
                        blocks.append((t0, "i", (im, w)))
                    except Exception as e:
                        log.warning("image crop failed: %s", e)
            blocks.sort(key=lambda b: b[0])

            for _, kind, data in blocks:
                if kind == "t":
                    t = doc.add_table(rows=len(data), cols=len(data[0]))
                    t.style = "Table Grid"
                    for r, row in enumerate(data):
                        for c, val in enumerate(row):
                            t.cell(r, c).text = val
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_after = Pt(2)
                    stats["tables"] += 1
                elif kind == "p":
                    p = data
                    para = doc.add_paragraph()
                    pf = para.paragraph_format
                    pf.space_after = Pt(0)
                    pf.line_spacing = 1.0
                    if p["gap"] is not None:
                        extra = max(0.0, p["gap"] - 0.3 * p["line_h"])
                        pf.space_before = Pt(min(extra * 0.9, 36))
                    centred = abs((p["x0"] + p["x1"]) / 2 - pw / 2) < 0.04 * pw and p["x0"] - p["left_edge"] > 0.08 * pw
                    if centred:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        ind = (p["x0"] - p["left_edge"]) / 72.0
                        if ind > 0.2:
                            pf.left_indent = Inches(min(ind, avail_w * 0.6))
                    run = para.add_run(p["text"])
                    run.font.size = Pt(max(6, min(p["size"], 60)))
                    run.bold = p["bold"] or None
                    run.italic = p["italic"] or None
                    stats["paragraphs"] += 1
                else:
                    im, wpt = data
                    buf = io.BytesIO()
                    im.save(buf, "PNG")
                    doc.add_picture(io.BytesIO(buf.getvalue()), width=Inches(min(wpt / 72.0, avail_w)))
                    stats["pictures"] += 1
            _free(page)
            if progress:
                progress(k + 1, len(idxs), f"Page {i + 1}")

    _chk(cancel)
    out_path = unique_path(out_path)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    tmp = out_path + ".part"
    doc.save(tmp)
    os.replace(tmp, out_path)
    res.outputs.append(out_path)
    res.stats = {"pages": len(idxs), **dict(stats)}
    res.notes.append("Layout is re-created from the PDF text, tables and pictures; check complex pages in Word.")
    return res


# ------------------------------------------------------------------ PDF -> Excel
_NUM_PLAIN = re.compile(r"^[+-]?\d+(\.\d+)?$")
_NUM_COMMA = re.compile(r"^[+-]?\d{1,3}(,\d{2,3})+(\.\d+)?$")


def to_number(s: str):
    """Convert clean numbers only. IDs, phone numbers, Aadhaar/PAN/account numbers, '007', long digit strings stay text."""
    t = s.strip()
    if not t:
        return s
    if _NUM_COMMA.match(t):
        t = t.replace(",", "")
    elif not _NUM_PLAIN.match(t):
        return s
    digits = re.sub(r"\D", "", t)
    if len(digits) > 15 or ("." not in t and len(digits) >= 10):
        return s
    body = t.lstrip("+-")
    if len(body) > 1 and body.startswith("0") and not body.startswith("0."):
        return s
    try:
        return float(t) if "." in t else int(t)
    except ValueError:
        return s


def _aligned_tables(page):
    """Borderless tables from text alignment. Guarded so ordinary text is not turned into a table."""
    try:
        ts = {"vertical_strategy": "text", "horizontal_strategy": "text", "snap_tolerance": 3,
              "join_tolerance": 3, "intersection_tolerance": 5, "min_words_vertical": 2, "min_words_horizontal": 2}
        tables = page.find_tables(table_settings=ts)
    except Exception:
        return []
    out = []
    for t in tables:
        rows = _clean_rows(t.extract())
        if len(rows) < 3 or not (2 <= len(rows[0]) <= 30):
            continue
        multi = sum(1 for r in rows if sum(1 for c in r if c) >= 2)
        if multi / len(rows) >= 0.7:
            out.append((t.bbox, rows))
    return out


def _text_columns(page):
    """Last resort: split text lines at wide gaps into columns."""
    txt = page.extract_text(layout=True) or ""
    rows = []
    for line in txt.splitlines():
        if line.strip():
            rows.append([c.strip() for c in re.split(r"\s{2,}", line.strip())])
    multi = sum(1 for r in rows if len(r) >= 2)
    if len(rows) >= 3 and multi / len(rows) >= 0.6:
        return rows, "columns"
    return [[r[0]] if len(r) == 1 else [" ".join(r)] for r in rows], "lines"


def _sheet_name(wb, wanted: str) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", wanted)[:31] or "Sheet"
    name, i = base, 2
    while name in wb.sheetnames:
        suffix = f"_{i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    return name


def pdf_to_excel(pdf: str, out_path: str, pages: str = "", password: Optional[str] = None,
                 layout: str = "sheet_per_table", text_fallback: bool = True, convert_numbers: bool = True,
                 progress=None, cancel=None) -> ToolResult:
    """layout: sheet_per_table | one_sheet | sheet_per_page"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    res = ToolResult(source=pdf)
    n = page_count(pdf, password)
    idxs = parse_pages(pages, n)
    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="D9E2F3")
    shared = {"ws": None, "row": 1}
    n_tables = n_rows = 0

    def write(ws, start_row, rows, header, title=None):
        r = start_row
        if title:
            ws.cell(r, 1, title).font = Font(bold=True, color="1F4E79")
            r += 1
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row, 1):
                v = to_number(val) if convert_numbers else val
                cell = ws.cell(r, ci, v)
                if isinstance(v, str) and "\n" in v:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if header and ri == 0:
                    cell.font, cell.fill = head_font, head_fill
            r += 1
        return r

    with _plumber(pdf, password) as plumber:
        for k, i in enumerate(idxs):
            _chk(cancel)
            page = plumber.pages[i]
            nchar = len((page.extract_text() or "").strip())
            if _is_scan(page, nchar) or nchar < 10:
                res.warnings.append(f"Page {i + 1}: no text layer (scanned image) — skipped. OCR is not included.")
                _free(page)
                if progress:
                    progress(k + 1, len(idxs), f"Page {i + 1}")
                continue
            found = [(rows, "ruled") for _, rows in _ruled_tables(page)]
            if not found:
                found = [(rows, "aligned") for _, rows in _aligned_tables(page)]
                if found:
                    res.warnings.append(f"Page {i + 1}: table detected from text alignment (no ruling lines) — please verify.")
            if not found and text_fallback:
                rows, kind = _text_columns(page)
                if rows:
                    found = [(rows, kind)]
                    res.warnings.append(f"Page {i + 1}: no table found — text exported as {'columns' if kind == 'columns' else 'one line per row'}.")
            for t_idx, (rows, kind) in enumerate(found, 1):
                header = kind in ("ruled", "aligned")
                label = f"Page {i + 1} – table {t_idx}" if kind in ("ruled", "aligned") else f"Page {i + 1} – text"
                if layout == "one_sheet":
                    if shared["ws"] is None:
                        shared["ws"] = wb.create_sheet("All tables")
                    shared["row"] = write(shared["ws"], shared["row"], rows, header, label) + 1
                    ws = shared["ws"]
                elif layout == "sheet_per_page":
                    nm = f"Page {i + 1}"
                    ws = wb[nm] if nm in wb.sheetnames else wb.create_sheet(_sheet_name(wb, nm))
                    start = ws.max_row + 2 if ws.max_row > 1 or ws.cell(1, 1).value else 1
                    write(ws, start, rows, header, label if len(found) > 1 else None)
                else:
                    nm = f"P{i + 1}_T{t_idx}" if header else f"P{i + 1}_text"
                    ws = wb.create_sheet(_sheet_name(wb, nm))
                    write(ws, 1, rows, header)
                    if header:
                        ws.freeze_panes = "A2"
                n_tables += 1
                n_rows += len(rows)
            _free(page)
            if progress:
                progress(k + 1, len(idxs), f"Page {i + 1}")

    if not wb.sheetnames:
        raise PdfToolError("No tables or text were found to export. If this PDF is a scan, it has no text layer "
                           "(OCR is not included in this version).")
    for ws in wb.worksheets:                                   # tidy column widths
        for col in ws.columns:
            width = max((len(str(c.value).split("\n")[0]) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 60)

    _chk(cancel)
    out_path = unique_path(out_path)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    tmp = out_path + ".part"
    wb.save(tmp)
    os.replace(tmp, out_path)
    res.outputs.append(out_path)
    res.stats = {"pages": len(idxs), "tables": n_tables, "rows": n_rows, "sheets": len(wb.sheetnames)}
    if convert_numbers:
        res.notes.append("Plain numbers became Excel numbers; IDs, phone/Aadhaar/account numbers and values like 007 stay text.")
    return res


# ------------------------------------------------------------------ batch driver used by the GUI
KINDS = ("images", "word", "excel")


def convert_many(kind: str, files: list, out_dir: Optional[str], opts: dict,
                 progress: Optional[Callable[[int, int, str], None]] = None,
                 cancel: Optional[Callable[[], bool]] = None) -> list:
    """Convert several PDFs. One bad file never stops the others. out_dir=None -> next to each PDF."""
    if kind not in KINDS:
        raise ValueError(kind)
    passwords = opts.get("passwords", {})
    results, plan, total = [], [], 0
    for f in files:
        try:
            n = page_count(f, passwords.get(f))
            idxs = parse_pages(opts.get("pages", ""), n)
            plan.append((f, len(idxs)))
            total += len(idxs)
        except Exception as e:
            results.append(ToolResult(source=f, error=str(e)))
    done = 0
    for fi, (f, cnt) in enumerate(plan):
        _chk(cancel)
        folder = out_dir or os.path.dirname(os.path.abspath(f))
        stem = os.path.splitext(os.path.basename(f))[0]
        pw = passwords.get(f)

        def sub(a, b, m, _d=done, _fi=fi):
            if progress:
                progress(_d + a, max(total, 1), f"[{_fi + 1}/{len(plan)}] {m}")
        try:
            if kind == "images":
                r = pdf_to_images(f, folder, opts.get("fmt", "png"), opts.get("dpi", 200), opts.get("pages", ""), pw,
                                  opts.get("mode", "render"), opts.get("jpeg_quality", 95), sub, cancel)
            elif kind == "word":
                r = pdf_to_word(f, os.path.join(folder, stem + ".docx"), opts.get("pages", ""), pw,
                                opts.get("keep_images", True), sub, cancel)
            else:
                r = pdf_to_excel(f, os.path.join(folder, stem + ".xlsx"), opts.get("pages", ""), pw,
                                 opts.get("layout", "sheet_per_table"), opts.get("text_fallback", True),
                                 opts.get("convert_numbers", True), sub, cancel)
        except Cancelled:
            raise
        except Exception as e:
            log.exception("PDF tool failed for %s", f)
            r = ToolResult(source=f, error=str(e) or e.__class__.__name__)
        results.append(r)
        done += cnt
    return results
