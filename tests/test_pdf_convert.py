import os

import pytest
from PIL import Image

pymupdf = pytest.importorskip("pymupdf")
reportlab = pytest.importorskip("reportlab")

from app.core import pdf_convert as C
from app.core.pdf_convert import ConvertOptions, smart_value
from app.core.pdf_engine import Cancelled, build_pdf


# ----------------------------------------------------------------- fixtures
@pytest.fixture(scope="module")
def pdfs(tmp_path_factory):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    d = tmp_path_factory.mktemp("pdfs")
    st = getSampleStyleSheet()
    text = str(d / "text.pdf")
    data = [["Item", "Qty", "Price"], ["Notebook", "10", "1,250.50"], ["Pen", "200", "0.75"],
            ["Aadhaar", "0123", "9876543210"], ["Refund", "(1,200)", "12%"]]
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(text, pagesize=A4).build([
        Paragraph("Invoice Report", st["Title"]),
        Paragraph("This is a normal text PDF with <b>bold</b> words and a table below.", st["Normal"]),
        Spacer(1, 12), t, Spacer(1, 12), Paragraph("Thank you for your business.", st["Normal"]),
        PageBreak(), Paragraph("Second page heading", st["Heading1"]), Paragraph("More text on page two.", st["Normal"])])
    # scanned copy: pages rendered to pictures only
    src = pymupdf.open(text)
    scan = pymupdf.open()
    for p in src:
        p.get_pixmap(dpi=150).save(str(d / "pg.png"))
        np_ = scan.new_page(width=p.rect.width, height=p.rect.height)
        np_.insert_image(np_.rect, filename=str(d / "pg.png"))
    scan_path = str(d / "scan.pdf")
    scan.save(scan_path)
    enc = str(d / "locked.pdf")
    src.save(enc, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    return {"text": text, "scan": scan_path, "locked": enc, "dir": d}


# ----------------------------------------------------------------- helpers
def test_parse_page_range():
    assert C.parse_page_range("", 5) == [0, 1, 2, 3, 4]
    assert C.parse_page_range("1-3,5", 6) == [0, 1, 2, 4]
    assert C.parse_page_range("4-", 6) == [3, 4, 5]
    assert C.parse_page_range("-2", 6) == [0, 1]
    assert C.parse_page_range("2,2,1", 6) == [0, 1]
    for bad in ("0", "abc", "5-3", "9"):
        with pytest.raises(ValueError):
            C.parse_page_range(bad, 6)


def test_unique_path_never_overwrites(tmp_path):
    p = tmp_path / "a.docx"
    p.write_text("x")
    assert C.unique_path(str(p)).endswith("a (1).docx")
    (tmp_path / "a (1).docx").write_text("x")
    assert C.unique_path(str(p)).endswith("a (2).docx")


@pytest.mark.parametrize("text,expected", [
    ("1,250.50", (1250.5, "#,##0.00")), ("12,34,567", (1234567, "#,##0")), ("0.75", (0.75, "0.00")),
    ("(1,200)", (-1200, "#,##0")), ("-5", (-5, "0")), ("12%", (0.12, "0%")), ("₹1,250", (1250, '"₹"#,##0')),
    ("Rs. 99.5", (99.5, '"₹"0.0')),
    ("0123", ("0123", None)), ("9876543210", ("9876543210", None)), ("123456789012", ("123456789012", None)),
    ("12/05/2024", ("12/05/2024", None)), ("A12", ("A12", None)), ("1.2.3", ("1.2.3", None)), ("", ("", None)),
])
def test_smart_value_keeps_identifiers_as_text(text, expected):
    assert smart_value(text) == expected


# ----------------------------------------------------------------- PDF -> Image
@pytest.mark.parametrize("fmt,ext", [("png", ".png"), ("jpg", ".jpg"), ("tiff", ".tiff")])
def test_images_render_formats_and_dpi(pdfs, tmp_path, fmt, ext):
    rep = C.pdf_to_images(pdfs["text"], str(tmp_path), ConvertOptions(fmt=fmt, dpi=100))
    assert rep.ok and len(rep.outputs) == 2 and all(o.endswith(ext) for o in rep.outputs)
    w, h = Image.open(rep.outputs[0]).size
    assert abs(w - 827) <= 2 and abs(h - 1169) <= 2          # A4 at 100 DPI


def test_images_page_range_and_no_overwrite(pdfs, tmp_path):
    a = C.pdf_to_images(pdfs["text"], str(tmp_path), ConvertOptions(pages="2"))
    b = C.pdf_to_images(pdfs["text"], str(tmp_path), ConvertOptions(pages="2"))
    assert len(a.outputs) == 1 and "page_002" in a.outputs[0]
    assert os.path.dirname(a.outputs[0]) != os.path.dirname(b.outputs[0])          # second run -> new folder


def test_images_extract_keeps_original_jpeg_bytes(tmp_path):
    import numpy as np
    src = str(tmp_path / "photo.jpg")
    arr = (np.random.default_rng(5).random((300, 400, 3)) * 255).astype("uint8")
    Image.fromarray(arr).save(src, "JPEG", quality=88)
    pdf = str(tmp_path / "p.pdf")
    build_pdf([src], pdf)
    rep = C.pdf_to_images(pdf, str(tmp_path), ConvertOptions(image_mode="extract"))
    assert rep.ok and len(rep.outputs) == 1
    assert open(rep.outputs[0], "rb").read() == open(src, "rb").read()               # bit-identical


def test_images_extract_reports_when_no_images(pdfs, tmp_path):
    rep = C.pdf_to_images(pdfs["text"], str(tmp_path), ConvertOptions(image_mode="extract"))
    assert not rep.outputs and any("No embedded images" in w for w in rep.warnings)


# ----------------------------------------------------------------- PDF -> Word
def test_word_from_text_pdf(pdfs, tmp_path):
    import docx
    rep = C.pdf_to_word(pdfs["text"], str(tmp_path / "o.docx"), ConvertOptions())
    assert rep.ok and not rep.ocr_pages
    assert rep.text_coverage is not None and rep.text_coverage >= 0.98
    d = docx.Document(rep.outputs[0])
    body = " ".join(p.text for p in d.paragraphs) + " ".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    for word in ("Invoice Report", "Notebook", "1,250.50", "Second page heading", "Thank you"):
        assert word in body, word
    assert len(d.tables) >= 1


def test_word_page_range(pdfs, tmp_path):
    import docx
    rep = C.pdf_to_word(pdfs["text"], str(tmp_path / "o.docx"), ConvertOptions(pages="2"))
    body = " ".join(p.text for p in docx.Document(rep.outputs[0]).paragraphs)
    assert "Second page heading" in body and "Invoice Report" not in body


def test_word_scanned_without_ocr_keeps_page_as_picture(pdfs, tmp_path):
    rep = C.pdf_to_word(pdfs["scan"], str(tmp_path / "o.docx"), ConvertOptions(ocr="off"))
    assert rep.ok and any("scanned" in w for w in rep.warnings)


@pytest.mark.skipif(not C.ocr_available(), reason="OCR component not installed")
def test_word_scanned_with_ocr(pdfs, tmp_path):
    import docx
    rep = C.pdf_to_word(pdfs["scan"], str(tmp_path / "o.docx"), ConvertOptions())
    assert rep.ok and rep.ocr_pages == [1, 2]
    body = " ".join(p.text for p in docx.Document(rep.outputs[0]).paragraphs)
    assert "Invoice Report" in body and "Notebook" in body and "Second page heading" in body
    assert any("OCR text can contain mistakes" in w for w in rep.warnings)


# ----------------------------------------------------------------- PDF -> Excel
def test_excel_tables_numbers_and_identifiers(pdfs, tmp_path):
    from openpyxl import load_workbook
    rep = C.pdf_to_excel(pdfs["text"], str(tmp_path / "o.xlsx"), ConvertOptions(pages="1"))
    assert rep.ok and rep.tables == 1
    ws = load_workbook(rep.outputs[0]).worksheets[0]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["Item", "Qty", "Price"]
    assert rows[1] == ["Notebook", 10, 1250.5]
    assert rows[2] == ["Pen", 200, 0.75]
    assert rows[3] == ["Aadhaar", "0123", "9876543210"]            # identifiers remain text
    assert rows[4] == ["Refund", -1200, 0.12]
    assert ws["C2"].number_format == "#,##0.00"


def test_excel_numbers_can_stay_text(pdfs, tmp_path):
    from openpyxl import load_workbook
    rep = C.pdf_to_excel(pdfs["text"], str(tmp_path / "o.xlsx"), ConvertOptions(pages="1", convert_numbers=False))
    ws = load_workbook(rep.outputs[0]).worksheets[0]
    assert ws["C2"].value == "1,250.50" and ws["B2"].value == "10"


@pytest.mark.parametrize("layout,sheets", [("sheet_per_table", 2), ("sheet_per_page", 2), ("single_sheet", 1)])
def test_excel_layouts(pdfs, tmp_path, layout, sheets):
    from openpyxl import load_workbook
    rep = C.pdf_to_excel(pdfs["text"], str(tmp_path / "o.xlsx"), ConvertOptions(excel_layout=layout))
    assert rep.ok and len(load_workbook(rep.outputs[0]).worksheets) == sheets


def test_excel_page_without_tables_uses_text_fallback_with_warning(pdfs, tmp_path):
    rep = C.pdf_to_excel(pdfs["text"], str(tmp_path / "o.xlsx"), ConvertOptions(pages="2"))
    assert rep.ok and any("no table lines detected" in w for w in rep.warnings)
    rep2 = C.pdf_to_excel(pdfs["text"], str(tmp_path / "o2.xlsx"), ConvertOptions(pages="2", text_fallback=False))
    assert not rep2.ok and "No tables or text" in rep2.error


@pytest.mark.skipif(not C.ocr_available(), reason="OCR component not installed")
def test_excel_scanned_with_ocr(pdfs, tmp_path):
    from openpyxl import load_workbook
    rep = C.pdf_to_excel(pdfs["scan"], str(tmp_path / "o.xlsx"), ConvertOptions(pages="1"))
    assert rep.ok and rep.ocr_pages == [1]
    cells = [c.value for ws in load_workbook(rep.outputs[0]).worksheets for r in ws.iter_rows() for c in r if c.value is not None]
    assert "Notebook" in cells and any(v in (10, "10") for v in cells)


# ----------------------------------------------------------------- batch / errors / cancel
def test_password_protected_pdf(pdfs, tmp_path):
    bad = C.run_batch("word", [pdfs["locked"]], str(tmp_path), ConvertOptions())[0]
    assert not bad.ok and "password" in bad.error.lower()
    good = C.run_batch("word", [pdfs["locked"]], str(tmp_path), ConvertOptions(password="secret"))[0]
    assert good.ok


def test_one_bad_file_does_not_stop_the_batch(pdfs, tmp_path):
    junk = tmp_path / "junk.pdf"
    junk.write_bytes(b"this is not a pdf")
    notpdf = tmp_path / "pic.pdf"
    Image.new("RGB", (50, 50)).save(str(tmp_path / "pic.png"))
    os.rename(str(tmp_path / "pic.png"), str(notpdf))
    reps = C.run_batch("excel", [str(junk), str(notpdf), pdfs["text"]], str(tmp_path / "out"), ConvertOptions(pages="1"))
    assert [r.ok for r in reps] == [False, False, True]
    assert all(r.error for r in reps[:2])


def test_cancel_and_source_untouched(pdfs, tmp_path):
    before = open(pdfs["text"], "rb").read()
    with pytest.raises(Cancelled):
        C.run_batch("image", [pdfs["text"]], str(tmp_path), ConvertOptions(), cancel=lambda: True)
    C.run_batch("word", [pdfs["text"]], str(tmp_path), ConvertOptions())
    assert open(pdfs["text"], "rb").read() == before


def test_invalid_page_range_is_reported_not_crashed(pdfs, tmp_path):
    rep = C.run_batch("image", [pdfs["text"]], str(tmp_path), ConvertOptions(pages="9"))[0]
    assert not rep.ok and "beyond" in rep.error


def test_report_text_is_readable(pdfs, tmp_path):
    rep = C.run_batch("excel", [pdfs["text"]], str(tmp_path), ConvertOptions(pages="1"))[0]
    t = rep.to_text()
    assert "PDF → Excel" in t and "Tables: 1" in t
